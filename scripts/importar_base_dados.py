"""
Importa a planilha "Base_de_dados.xlsx" (formato simplificado, com alunos
fictícios para o projeto integrador) para as tabelas `students` e
`aee_info` — tudo em um único arquivo/planilha, ao contrário do formato
PILOTO + AEE separado usado por import_students.py/import_aee.py.

Colunas esperadas (primeira aba da planilha):
    Turno | Turma | Nome do Aluno | Filiação 1 | Data Nascimento | RA |
    Código | Data Matrícula | AEE

A coluna "AEE" pode vir vazia (aluno sem AEE) ou com o texto da condição
(ex: "TEA", "CADEIRANTE", "DEF. INTELECTUAL", ...) — qualquer valor
preenchido marca o aluno com aee=1 e grava esse texto em aee_info.deficiencia.

A série (1º ao 9º Ano) é deduzida automaticamente do começo do código da
turma (ex: "1A" -> "1º Ano", "9C" -> "9º Ano").

Uso:
    python scripts/importar_base_dados.py data/raw/Base_de_dados.xlsx
"""
import argparse
import re
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

from models import get_db, init_db, agora, salvar_aee_info

COLUNAS = [
    "turno", "turma", "nome", "filiacao1", "data_nascimento",
    "ra", "codigo", "data_matricula", "aee",
]


def _to_iso(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    return None


def _serie_da_turma(turma):
    m = re.match(r"^(\d+)", str(turma).strip())
    return f"{m.group(1)}º Ano" if m else None


def carregar_planilha(path):
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    linhas = []
    for r in rows[1:]:
        if not r or not r[2]:  # sem nome
            continue
        linhas.append(dict(zip(COLUNAS, r)))
    return linhas


def importar(path, dry_run=False):
    linhas = carregar_planilha(path)
    print(f"Linhas na planilha: {len(linhas)}")

    if dry_run:
        for l in linhas[:5]:
            print(" ", l["nome"], l["turma"], l["aee"] or "-")
        return

    init_db()
    novos, atualizados, aee_marcados = 0, 0, 0
    with get_db() as db:
        for l in linhas:
            codigo = str(l["codigo"]).strip()
            turma = str(l["turma"]).strip()
            serie = _serie_da_turma(turma)
            aee_texto = (str(l["aee"]).strip() if l["aee"] else None)

            existente = db.execute(
                "SELECT id FROM students WHERE codigo_interno = ?", (codigo,)
            ).fetchone()
            valores = (
                str(l["ra"]) if l["ra"] else None,
                l["nome"].strip(),
                l["filiacao1"],
                _to_iso(l["data_nascimento"]),
                turma,
                serie,
                "ATIVO",
                1 if aee_texto else 0,
            )
            if existente:
                student_id = existente["id"]
                db.execute(
                    """UPDATE students SET ra_prodesp=?, nome=?, filiacao1=?, data_nascimento=?,
                       turma=?, serie=?, situacao=?, aee=?, updated_at=?
                       WHERE codigo_interno=?""",
                    valores + (agora(), codigo),
                )
                atualizados += 1
            else:
                cur = db.execute(
                    """INSERT INTO students
                       (codigo_interno, ra_prodesp, nome, filiacao1, data_nascimento,
                        turma, serie, situacao, aee)
                       VALUES (?,?,?,?,?,?,?,?,?)""",
                    (codigo,) + valores,
                )
                student_id = cur.lastrowid
                novos += 1

            if aee_texto:
                salvar_aee_info(
                    db, student_id, eol=codigo, laudo=None,
                    observacoes="Importado da base fictícia do projeto integrador.",
                    deficiencia=aee_texto, turno=l["turno"],
                )
                aee_marcados += 1

    print(f"Importação concluída: {novos} aluno(s) novo(s), {atualizados} atualizado(s), "
          f"{aee_marcados} marcado(s) como AEE.")
    return {"total": len(linhas), "novos": novos, "atualizados": atualizados, "aee_marcados": aee_marcados}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("planilha", help="Caminho do .xlsx (Base_de_dados.xlsx)")
    parser.add_argument("--dry-run", action="store_true", help="Só mostra o que seria importado.")
    args = parser.parse_args()
    importar(args.planilha, args.dry_run)
