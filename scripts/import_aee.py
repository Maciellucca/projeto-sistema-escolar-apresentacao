"""
Importa a planilha de alunos AEE (Atendimento Educacional Especializado) e:
  1. marca `students.aee = 1` para os alunos encontrados (usado para o
     badge "AEE" visível no front-end para qualquer profissional);
  2. grava o detalhe (laudo/diagnóstico) em `aee_info`, tabela cujo
     conteúdo só é exibido no front-end para os perfis Admin/AEE.

O cruzamento é feito pelo código EOL da planilha, que corresponde ao
"Código" interno do aluno na planilha PILOTO — portanto rode
import_students.py ANTES deste script.

Uso:
    python scripts/import_aee.py data/raw/AEE_S_2026.xlsx
"""
import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

from models import get_db, init_db, agora, salvar_aee_info


def importar(path, dry_run=False):
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h else "" for h in rows[0]]

    encontrados, nao_encontrados = 0, []

    init_db()
    with get_db() as db:
        for r in rows[1:]:
            if not r or not r[1]:
                continue
            nome, eol = r[1], r[2]
            turno = r[4]
            laudo = r[5]
            turma_obs = r[6]
            projeto_rede = r[7]
            deficiencia = r[8]

            if not eol:
                nao_encontrados.append((nome, "sem código EOL na planilha"))
                continue

            aluno = db.execute(
                "SELECT id FROM students WHERE codigo_interno = ?", (str(eol),)
            ).fetchone()
            if not aluno:
                nao_encontrados.append((nome, f"EOL {eol} não encontrado nos alunos importados"))
                continue

            encontrados += 1
            if dry_run:
                continue

            student_id = aluno["id"]
            db.execute(
                "UPDATE students SET aee = 1, updated_at = ? WHERE id = ?",
                (agora(), student_id),
            )
            salvar_aee_info(db, student_id, str(eol), laudo, projeto_rede, deficiencia, turno)

    print(f"AEE cruzados com sucesso: {encontrados}")
    if nao_encontrados:
        print(f"Não localizados na base de alunos importada ({len(nao_encontrados)}):")
        for nome, motivo in nao_encontrados[:30]:
            print(f"   - {nome}: {motivo}")
        print("Isso é esperado para alunos de outras séries/turmas que ainda não foram importadas.")
    return {"encontrados": encontrados, "nao_encontrados": nao_encontrados}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("planilha", help="Caminho do .xlsx da lista AEE")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    importar(args.planilha, args.dry_run)
