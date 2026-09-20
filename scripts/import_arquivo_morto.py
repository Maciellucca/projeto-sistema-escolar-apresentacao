"""
Importa o "arquivo morto" — registro permanente de todo aluno que já
passou pela escola, com uma aba por letra do alfabeto (A a Z), cada
uma com as colunas Número, Nome e Data de Nascimento (sem cabeçalho —
os dados começam direto na linha 1).

Linhas com número preenchido mas nome vazio são ignoradas (o arquivo
original tem o número pré-preenchido até o fim de cada aba, mas só
uma parte das linhas tem aluno de verdade cadastrado).

A importação é idempotente: rodar de novo com a mesma planilha (ou uma
atualizada) não duplica quem já foi importado (confere por
letra+número+nome antes de inserir).

Formato aceito: .xlsx (se você tiver um .xls antigo, salve como .xlsx
no Excel/LibreOffice antes de importar — "Salvar como" → "Pasta de
trabalho do Excel (.xlsx)").

Uso:
    python scripts/import_arquivo_morto.py planilha.xlsx
"""
import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

from models import get_db, init_db


def _limpo(valor):
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def _data_iso(valor):
    if valor is None:
        return None
    if hasattr(valor, "strftime"):
        return valor.strftime("%Y-%m-%d")
    return _limpo(valor)


def importar(caminho):
    init_db()
    wb = load_workbook(caminho, data_only=True)

    resultado = {"importados": 0, "duplicados_ignorados": 0, "abas_processadas": [], "abas_ignoradas": []}

    with get_db() as db:
        for nome_aba in wb.sheetnames:
            letra = nome_aba.strip().upper()
            if len(letra) != 1 or not letra.isalpha():
                resultado["abas_ignoradas"].append(nome_aba)
                continue

            ws = wb[nome_aba]
            resultado["abas_processadas"].append(letra)

            for linha in ws.iter_rows(values_only=True):
                if not linha or len(linha) < 2:
                    continue
                nome = _limpo(linha[1])
                if not nome:
                    continue
                numero = _limpo(linha[0])
                data_nasc = _data_iso(linha[2]) if len(linha) > 2 else None

                existente = db.execute(
                    "SELECT id FROM arquivo_morto WHERE letra=? AND numero=? AND nome=?",
                    (letra, numero, nome),
                ).fetchone()
                if existente:
                    resultado["duplicados_ignorados"] += 1
                    continue

                db.execute(
                    """INSERT INTO arquivo_morto (letra, numero, nome, data_nascimento, origem)
                       VALUES (?,?,?,?,'importado')""",
                    (letra, numero, nome, data_nasc),
                )
                resultado["importados"] += 1

    return resultado


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("planilha")
    args = parser.parse_args()
    r = importar(args.planilha)

    print(f"Abas processadas: {', '.join(r['abas_processadas']) or '(nenhuma)'}")
    if r["abas_ignoradas"]:
        print(f"Abas ignoradas (nome não é uma letra só): {', '.join(r['abas_ignoradas'])}")
    print(f"Registros importados: {r['importados']}")
    if r["duplicados_ignorados"]:
        print(f"Registros já existentes (ignorados, sem duplicar): {r['duplicados_ignorados']}")
